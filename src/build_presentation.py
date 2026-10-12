"""
build_presentation.py - generates the college PowerPoint (14 slides)
====================================================================
Every number and chart comes from the REAL project outputs:
  metrics : output/plots/model_comparison.csv  (written by train_and_evaluate.py)
  charts  : outputs/*.png (EDA) and output/plots/*.png (results)
Re-run training first if you change the models, then:
    ./.venv/bin/python src/build_presentation.py
"""

import os

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(PROJECT_ROOT, "docs", "Loan_Approval_Presentation.pptx")
PLOTS = os.path.join(PROJECT_ROOT, "output", "plots")
EDA = os.path.join(PROJECT_ROOT, "outputs")

METRICS = pd.read_csv(os.path.join(PLOTS, "model_comparison.csv"), index_col=0)
BEST = METRICS["CV F1"].idxmax()
MODELS = list(METRICS.index)

# ---- design system: deep "banking" green + gold accent --------------------
GREEN = RGBColor(0x0F, 0x3D, 0x3E)
GREEN_2 = RGBColor(0x1F, 0x6F, 0x5C)
GOLD = RGBColor(0xE3, 0xA7, 0x2F)
MINT = RGBColor(0xEA, 0xF3, 0xF0)
INK = RGBColor(0x1E, 0x2B, 0x2B)
MUTED = RGBColor(0x5B, 0x6B, 0x6B)
RED = RGBColor(0xB5, 0x44, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD_FONT, BODY_FONT = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 14
HEADINGS = {}   # slide_id -> title, used for the markdown outline


# ===========================================================================
# helpers
# ===========================================================================
def new_slide(title=None, dark=False, heading=None):
    slide = prs.slides.add_slide(BLANK)
    HEADINGS[slide.slide_id] = title or heading
    if dark:
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = GREEN
    if title:
        text(slide, title, 0.6, 0.4, 12.1, 0.85, size=34, bold=True,
             color=GREEN, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE)
    return slide


def text(slide, content, x, y, w, h, size=16, bold=False, color=INK,
         font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         italic=False):
    """content = str, or a list of str / (str, {overrides}) paragraphs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        para_text, opts = (para if isinstance(para, tuple) else (para, {}))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space", 6))
        run = p.add_run()
        run.text = para_text
        f = run.font
        f.name = opts.get("font", font)
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic)
        f.color.rgb = opts.get("color", color)
    return box


def bullets(slide, items, x, y, w, h, size=16, color=INK):
    """Real bulleted paragraphs (round bullet via XML)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.28)))
        pPr.set("indent", str(-Inches(0.25)))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar",
                             {"char": "•"})
        pPr.append(bu)
        # support "**bold lead** rest" by a simple split on the first ':'
        if ": " in item and item.index(": ") < 40:
            lead, rest = item.split(": ", 1)
            parts = [(lead + ": ", True), (rest, False)]
        else:
            parts = [(item, False)]
        for t, b in parts:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), b, BODY_FONT
            r.font.color.rgb = color
    return box


def card(slide, x, y, w, h, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def badge(slide, x, y, label, fill=GOLD, color=GREEN, d=0.6, size=18):
    """Numbered circle - the deck's repeated visual motif."""
    c = card(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), True, color
    r.font.name = HEAD_FONT
    return c


def image(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Inches(x), Inches(y),
                                    width=Inches(w) if w else None,
                                    height=Inches(h) if h else None)


def finish(slide, number, notes, dark=False):
    text(slide, f"{number} / {TOTAL}", 11.9, 7.0, 0.9, 0.3, size=10,
         color=MINT if dark else MUTED, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = notes


def m(model, metric):
    return METRICS.loc[model, metric]


# ===========================================================================
# 1. Title
# ===========================================================================
s = new_slide(dark=True, heading="Loan Approval Prediction")
text(s, "AI COURSE PROJECT  ·  2ND YEAR", 0.9, 1.5, 11, 0.4, size=14,
     bold=True, color=GOLD)
text(s, "Loan Approval Prediction", 0.9, 2.0, 11.5, 1.2, size=54, bold=True,
     color=WHITE, font=HEAD_FONT)
text(s, "Predicting whether a loan will be approved, using machine learning "
        "on applicant data", 0.9, 3.2, 11, 0.8, size=22, color=MINT)
text(s, ["Presented by:  <Your Name>   ·   Roll No: <...>",
         "Guide: <...>   ·   <College name>   ·   <Date>"],
     0.9, 5.3, 11, 1.0, size=16, color=WHITE)
finish(s, 1, "Greet the panel. One line: 'My project predicts whether a loan "
             "application will be approved or rejected using machine learning "
             "on a public dataset of 614 applications.' Keep this slide short.",
       dark=True)

# ===========================================================================
# 2. Problem statement
# ===========================================================================
s = new_slide("The problem")
text(s, "Banks receive thousands of loan applications. Checking each one by "
        "hand is slow, inconsistent and subjective.", 0.6, 1.35, 12, 0.8,
     size=18, color=MUTED)
for i, (head, body, colr) in enumerate([
    ("Mistake 1: approve a risky applicant",
     "The customer defaults and the bank loses money (a false positive).", RED),
    ("Mistake 2: reject a good applicant",
     "The bank loses a good customer and the business (a false negative).", GOLD),
]):
    x = 0.6 + i * 6.15
    card(s, x, 2.45, 5.85, 1.9)
    text(s, head, x + 0.3, 2.6, 5.3, 0.5, size=19, bold=True, color=colr)
    text(s, body, x + 0.3, 3.15, 5.3, 1.1, size=16)
card(s, 0.6, 4.75, 12.0, 1.8, fill=GREEN)
text(s, "OUR GOAL", 0.95, 4.95, 3, 0.4, size=13, bold=True, color=GOLD)
text(s, "Learn from past applications and predict Approved / Rejected "
        "automatically, and explain which features drive the decision.",
     0.95, 5.35, 11.3, 1.0, size=20, color=WHITE)
finish(s, 2, "Stress the two kinds of mistakes. They cost different amounts, "
             "and that is why we report precision AND recall instead of just "
             "accuracy. Technically this is supervised binary classification "
             "on tabular data.")

# ===========================================================================
# 3. Dataset
# ===========================================================================
s = new_slide("The dataset")
stats = [("614", "loan applications"), ("13", "columns (11 used as inputs)"),
         ("68.7%", "approved: imbalanced classes"), ("149", "missing cells in 7 columns")]
for i, (num, label) in enumerate(stats):
    x = 0.6 + i * 3.08
    card(s, x, 1.45, 2.85, 1.85)
    text(s, num, x + 0.25, 1.55, 2.4, 1.0, size=44, bold=True, color=GREEN,
         font=HEAD_FONT)
    text(s, label, x + 0.25, 2.5, 2.45, 0.7, size=14, color=MUTED)
bullets(s, [
    "Source: Analytics Vidhya \"Loan Prediction\" practice problem (public; also on Kaggle).",
    "Data: fictional \"Dream Housing Finance\" sample data, not a real bank's records.",
    "Target: Loan_Status, Y = approved (422) / N = rejected (192).",
    "Inputs: 6 categorical (Gender, Married, Dependents, Education, Self_Employed, "
    "Property_Area) and 5 numeric (2 incomes, LoanAmount, Term, Credit_History).",
], 0.6, 3.75, 7.0, 3.1, size=16)
image(s, os.path.join(EDA, "1_approval_vs_rejection.png"), 7.95, 3.6, w=4.65)
finish(s, 3, "Say where the data is from before anyone asks: a public teaching "
             "dataset. The 69/31 imbalance means a model that always says "
             "'approve' already gets 69% accuracy, so accuracy alone is not "
             "enough.")

# ===========================================================================
# 4. EDA
# ===========================================================================
s = new_slide("What the data told us (EDA)")
image(s, os.path.join(EDA, "4_credit_history_vs_approval.png"), 0.6, 1.4, w=7.6)
card(s, 8.6, 1.4, 4.1, 2.4, fill=GREEN)
text(s, "80% vs 8%", 8.9, 1.55, 3.6, 1.0, size=44, bold=True, color=GOLD,
     font=HEAD_FONT)
text(s, "approval rate with vs without a credit history (10× difference)",
     8.9, 2.55, 3.6, 1.1, size=16, color=WHITE)
bullets(s, [
    "Income is very skewed: median ₹3,812, mean ₹5,403, max ₹81,000 (skew 6.54).",
    "Property area has a mild effect: Semiurban 77%, Urban 66%, Rural 61%.",
    "Income alone barely separates the classes, but income compared with the loan "
    "size does. This led to our engineered features.",
], 0.6, 4.55, 12.1, 2.4, size=16)
finish(s, 4, "Tell the story: credit history dominates. Skewed income means we "
             "fill missing values with the MEDIAN, not the mean. Raw income on "
             "its own is weak, which motivated the affordability features.")

# ===========================================================================
# 5. Raw features
# ===========================================================================
s = new_slide("Features used: the 11 raw inputs")
rows = [
    ("Feature", "Type", "Meaning", "Signal"),
    ("Credit_History", "1 / 0", "Repaid past loans on time?", "Very strong"),
    ("ApplicantIncome", "₹ / month", "Main applicant's income", "Medium"),
    ("CoapplicantIncome", "₹ / month", "Co-applicant's income (0 if none)", "Medium"),
    ("LoanAmount", "₹ thousands", "Amount requested", "Medium"),
    ("Loan_Amount_Term", "months", "Repayment period (mostly 360)", "Low alone"),
    ("Property_Area", "3 categories", "Urban / Semiurban / Rural", "Weak"),
    ("Married · Dependents", "categories", "Family situation", "≈ none"),
    ("Education · Self_Employed", "yes / no", "Graduate? Self-employed?", "≈ none"),
    ("Gender", "M / F", "Sensitive attribute", "≈ none"),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.4), Inches(12.1),
                         Inches(5.3)).table
widths = [3.4, 1.9, 4.8, 2.0]
for c, wdt in enumerate(widths):
    tbl.columns[c].width = Inches(wdt)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = GREEN if r == 0 else (MINT if r % 2 else WHITE)
        cell.margin_left = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.size, run.font.name = Pt(15), BODY_FONT
        run.font.bold = r == 0 or c == 0
        strong = val in ("Very strong",)
        run.font.color.rgb = WHITE if r == 0 else (GREEN_2 if strong else INK)
finish(s, 5, "Walk down the table quickly. Point out the last column: it comes "
             "from our measured permutation importance (slide 12), not "
             "guesswork. Loan_ID is dropped because a unique ID has no meaning "
             "and could only be memorised.")

# ===========================================================================
# 6. Engineered features
# ===========================================================================
s = new_slide("Features we created: affordability")
text(s, "A loan officer asks: \"Can this household afford the monthly "
        "instalment?\" So we built 5 new columns that answer it.",
     0.6, 1.3, 12, 0.7, size=17, color=MUTED)
eng = [
    ("TotalIncome", "Applicant + Co-applicant income", "The bank lends to the household"),
    ("EMI", "LoanAmount × 1000 ÷ Term", "Approx. monthly instalment in ₹"),
    ("BalanceIncome", "TotalIncome − EMI", "Money left after paying the EMI"),
    ("LoanToIncome", "LoanAmount × 1000 ÷ TotalIncome", "Loan size in months of income"),
    ("LoanAmountLog", "log(1 + LoanAmount)", "Tames very large loans"),
]
for i, (name, formula, why) in enumerate(eng):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 2.2 + row * 2.05
    card(s, x, y, 3.85, 1.8)
    badge(s, x + 0.25, y + 0.25, str(i + 1), d=0.5, size=16)
    text(s, name, x + 0.9, y + 0.22, 2.85, 0.5, size=18, bold=True, color=GREEN)
    text(s, formula, x + 0.25, y + 0.85, 3.45, 0.45, size=14, color=GREEN_2,
         font="Courier New")
    text(s, why, x + 0.25, y + 1.25, 3.45, 0.5, size=14, color=MUTED)
card(s, 8.8, 4.25, 3.9, 1.8, fill=GREEN)
text(s, [("Did it help?", {"bold": True, "color": GOLD, "size": 15}),
         ("LR test ROC-AUC 0.852 → 0.870. LoanToIncome & BalanceIncome are "
          "the Random Forest's #2 and #3 columns.", {"color": WHITE, "size": 14})],
     9.05, 4.4, 3.45, 1.6)
text(s, "Computed inside the pipeline (src/features.py), so the app sends "
        "only raw values.", 0.6, 6.45, 12, 0.4, size=13, color=MUTED, italic=True)
finish(s, 6, "This is our main modification. The formulas use only the "
             "applicant's own row and learn nothing from the data, so there is "
             "no leakage. Evidence: LR's ROC-AUC improved and two engineered "
             "features ranked right after Credit_History in the forest.")

# ===========================================================================
# 7. Pipeline
# ===========================================================================
s = new_slide("How the data flows: one leakage-safe pipeline")
steps = [("Raw row", "11 columns"), ("Engineer", "+5 features"),
         ("Impute", "median / mode"), ("Transform", "StandardScaler\nOneHotEncoder"),
         ("Model", "LR · DT · RF"), ("Output", "Approved?\n+ probability")]
for i, (head, sub) in enumerate(steps):
    x = 0.6 + i * 2.1
    card(s, x, 1.6, 1.8, 1.9, fill=GREEN if i in (0, 5) else MINT)
    badge(s, x + 0.6, 1.35, str(i + 1), d=0.55, size=16)
    text(s, head, x + 0.1, 2.05, 1.6, 0.5, size=16, bold=True,
         color=WHITE if i in (0, 5) else GREEN, align=PP_ALIGN.CENTER)
    text(s, sub.split("\n"), x + 0.1, 2.55, 1.6, 0.9, size=13,
         color=MINT if i in (0, 5) else MUTED, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        a = card(s, x + 1.84, 2.42, 0.22, 0.28, fill=GOLD, shape=MSO_SHAPE.RIGHT_ARROW)
bullets(s, [
    "Split first: 80% train (491 rows) / 20% test (123 rows), stratified, random_state = 42.",
    "Fit on train only: imputers, scaler and encoder learn their values from the "
    "491 training rows and are then applied unchanged to test rows. No data leakage.",
    "One object: preprocessing + model are saved together (pickle), so the app "
    "processes inputs exactly like training did.",
], 0.6, 4.0, 12.1, 2.8, size=17)
finish(s, 7, "Define data leakage in one line: 'information from the test set "
             "sneaking into training and inflating the score.' The Pipeline "
             "makes it impossible because .fit() only ever sees X_train.")

# ===========================================================================
# 8. Models
# ===========================================================================
s = new_slide("Three models compared")
model_info = [
    ("Logistic Regression", "Linear baseline",
     ["Weighted sum → sigmoid → probability",
      "Needs scaled features",
      "Easy to interpret (weights)"]),
    ("Decision Tree", "Rule-based (max depth 4)",
     ["Learns yes/no questions using Gini impurity",
      "Depth capped to avoid memorising",
      "Human-readable rules"]),
    ("Random Forest", "Ensemble of 300 trees (new)",
     ["Each tree sees random rows + features",
      "Trees vote: their errors average out",
      "Usually beats one tree"]),
]
for i, (name, tag, pts) in enumerate(model_info):
    x = 0.6 + i * 4.1
    is_best = name == BEST
    card(s, x, 1.45, 3.85, 4.1, fill=GREEN if is_best else MINT)
    text(s, name, x + 0.3, 1.65, 3.3, 0.55, size=21, bold=True,
         color=WHITE if is_best else GREEN, font=HEAD_FONT)
    text(s, tag, x + 0.3, 2.2, 3.3, 0.4, size=14, bold=True, color=GOLD)
    bullets(s, pts, x + 0.3, 2.8, 3.35, 2.4, size=15,
            color=WHITE if is_best else INK)
    if is_best:
        text(s, "★ Selected model", x + 0.3, 4.95, 3.3, 0.4, size=14,
             bold=True, color=GOLD)
text(s, "All three share the same features and preprocessing, so the "
        "comparison is fair.", 0.6, 5.9, 12, 0.4, size=14, color=MUTED,
     italic=True)
finish(s, 8, "Sigmoid: the S-curve mapping any score into 0 to 1, read as a "
             "probability. Gini: how mixed a node is; the tree picks splits "
             "that make nodes purer. Random Forest = bagging + random feature "
             "subsets. We added it as a modification.")

# ===========================================================================
# 9. Evaluation method
# ===========================================================================
s = new_slide("How we evaluate fairly")
card(s, 0.6, 1.45, 5.9, 5.2)
text(s, "5-fold cross-validation (new)", 0.9, 1.6, 5.4, 0.5, size=19,
     bold=True, color=GREEN)
for k in range(5):
    y = 2.3 + k * 0.5
    for j in range(5):
        fill = GOLD if j == k else GREEN_2
        card(s, 0.9 + j * 0.85, y, 0.78, 0.38, fill=fill, shape=MSO_SHAPE.RECTANGLE)
text(s, "Gold = validation fold, green = training folds. Five rounds on the "
        "491 training rows; the average F1 picks the winner. The test set is "
        "only used once, at the end.", 0.9, 4.9, 5.4, 1.6, size=14, color=MUTED)
metrics_info = [
    ("Accuracy", "% of all decisions correct (misleading with a 69/31 imbalance)"),
    ("Precision", "Of loans we approved, how many were truly good? Protects the bank"),
    ("Recall", "Of truly good loans, how many did we approve? Protects customers"),
    ("F1-score", "Balance of precision & recall (our selection metric)"),
    ("ROC-AUC", "How well applicants are ranked (0.5 = coin flip, 1 = perfect)"),
]
for i, (name, desc) in enumerate(metrics_info):
    y = 1.5 + i * 1.03
    badge(s, 6.9, y + 0.05, str(i + 1), d=0.55, size=17)
    text(s, name, 7.6, y - 0.02, 5.0, 0.4, size=17, bold=True, color=GREEN)
    text(s, desc, 7.6, y + 0.36, 5.1, 0.6, size=14, color=MUTED)
finish(s, 9, "Why cross-validation? One test set of 123 rows is a single lucky "
             "or unlucky draw. CV gives an average and a spread. We pick the "
             "model on CV F1 so the test set never influences the choice.")

# ===========================================================================
# 10. Results (native chart + table)
# ===========================================================================
s = new_slide("Results on the unseen test set")
CHART_METRICS = ["Accuracy", "Precision", "Recall", "F1-score", "ROC-AUC"]
cd = CategoryChartData()
cd.categories = CHART_METRICS
for name in MODELS:
    cd.add_series(name, [float(m(name, k)) for k in CHART_METRICS])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.35),
                        Inches(7.7), Inches(5.5), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(13)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0.6, 1.0
ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
ch.value_axis.tick_labels.font.size = Pt(12)
ch.category_axis.tick_labels.font.size = Pt(13)
for series, colr in zip(ch.series, [GREEN_2, GOLD, GREEN]):
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = colr
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.number_format = "0.00"
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(10)

card(s, 8.6, 1.45, 4.1, 2.3, fill=GREEN)
text(s, f"{m(BEST, 'Accuracy'):.1%}", 8.9, 1.55, 3.6, 1.0, size=44, bold=True,
     color=GOLD, font=HEAD_FONT)
text(s, f"test accuracy of the {BEST} (F1 {m(BEST, 'F1-score'):.3f}, "
        f"CV F1 {m(BEST, 'CV F1'):.3f} ± {m(BEST, 'CV F1 std'):.3f})",
     8.9, 2.55, 3.6, 1.1, size=14, color=WHITE)
bullets(s, [
    f"Random Forest is best on CV F1, test accuracy and test F1.",
    "The gap is small: all three CV F1 scores are within one standard deviation.",
    "No overfitting: test accuracy ≥ train accuracy for every model.",
], 8.6, 4.05, 4.1, 2.9, size=15)
finish(s, 10, "Quote the numbers from the chart. Be honest that the three "
              "models are close (CV F1 0.863 to 0.867), so the forest wins "
              "narrowly. Logistic Regression has the best test ROC-AUC (0.870) "
              "and is simpler, which is a fair alternative choice.")

# ===========================================================================
# 11. Confusion matrices
# ===========================================================================
s = new_slide("Where the models make mistakes")
image(s, os.path.join(PLOTS, "04_confusion_matrices.png"), 0.6, 1.4, w=12.1)
bullets(s, [
    "Random Forest: 84 of 85 good loans approved (only 1 missed), 24 of 38 risky loans rejected.",
    "Weak spot: about 1 in 3 risky applicants is still approved (14 false positives). "
    "The models are lenient, because the data has twice as many approvals.",
], 0.6, 5.5, 12.1, 1.4, size=16)
finish(s, 11, "Read one matrix aloud: rows = truth, columns = prediction. "
              "Admit the weakness yourself. The model rarely rejects a good "
              "customer but lets some risky ones through. Fix ideas: class "
              "weights or a higher decision threshold.")

# ===========================================================================
# 12. Feature importance
# ===========================================================================
s = new_slide("Which features really matter?")
image(s, os.path.join(PLOTS, "06_feature_importance.png"), 0.5, 1.35, w=7.4)
card(s, 8.3, 1.45, 4.4, 5.3)
text(s, "Permutation importance", 8.6, 1.6, 3.9, 0.5, size=18, bold=True,
     color=GREEN)
text(s, "Shuffle one column and see how much ROC-AUC drops. A big drop means "
        "the model relies on it.", 8.6, 2.1, 3.9, 1.0, size=14, color=MUTED)
bullets(s, [
    "Credit history: by far #1 (−0.29 AUC).",
    "Affordability next: co-applicant income, applicant income, loan amount.",
    "Gender, education, marital status, dependents: ≈ 0.",
    "The model decides on credit + affordability, not demographics.",
], 8.6, 3.2, 3.9, 3.5, size=15)
finish(s, 12, "This answers 'explain the features'. Shuffling a raw income "
              "column also scrambles the engineered features built from it, "
              "so each bar is the total value of that input. Fairness point: "
              "sensitive attributes contribute nothing and could be dropped.")

# ===========================================================================
# 13. App demo
# ===========================================================================
s = new_slide("Live demo: Streamlit web app")
for i, (head, body) in enumerate([
    ("Simple application form", "About you, income & credit, loan details"),
    ("Clear decision", "Likely / not likely to be approved, with plain reasons and the EMI"),
    ("Internals stay hidden", "Model names and probabilities appear only in an examiner view (?examiner=1)"),
]):
    y = 1.5 + i * 1.45
    badge(s, 0.6, y, str(i + 1), d=0.65, size=20)
    text(s, head, 1.5, y - 0.05, 5.0, 0.45, size=19, bold=True, color=GREEN)
    text(s, body, 1.5, y + 0.42, 5.0, 0.8, size=15, color=MUTED)
card(s, 7.0, 1.45, 5.7, 3.3, fill=GREEN)
text(s, [("RUN IT", {"bold": True, "color": GOLD, "size": 14}),
         ("streamlit run src/app.py", {"font": "Courier New", "color": WHITE,
                                       "size": 18, "bold": True, "space": 18}),
         ("Try this:", {"bold": True, "color": GOLD, "size": 14}),
         ("1. Keep the defaults → Likely to be approved (High)",
          {"color": WHITE, "size": 15}),
         ("2. Answer \"No\" to repaying past loans → Not likely (Very low)",
          {"color": WHITE, "size": 15}),
         ("Open ?examiner=1 to show all 3 models agree.",
          {"color": MINT, "size": 14, "italic": True})],
     7.35, 1.7, 5.1, 2.9)
text(s, "Academic demo only, not a real lending decision.", 0.6, 5.9, 12, 0.4,
     size=13, color=MUTED, italic=True)
finish(s, 13, "Run the app live if allowed. Flip the credit-history answer "
              "and show the decision change. Then open the examiner view "
              "(add ?examiner=1 to the URL) to show the three models' "
              "probabilities. Explain that applicants never see these: only "
              "a decision, simple reasons and their EMI.")

# ===========================================================================
# 14. Conclusion
# ===========================================================================
s = new_slide(dark=True, heading="Conclusion & future scope")
text(s, "Conclusion & future scope", 0.6, 0.45, 12, 0.85, size=34, bold=True,
     color=WHITE, font=HEAD_FONT)
cols = [
    ("What we achieved", [
        f"End-to-end ML system: {m(BEST, 'Accuracy'):.1%} test accuracy, "
        f"F1 {m(BEST, 'F1-score'):.3f}",
        "5 affordability features that measurably help",
        "Fair comparison of 3 models using cross-validation",
        "Working web app"]),
    ("Limitations", [
        "Small dataset (614), not from a real lender",
        "Lenient: ~1 in 3 risky applicants still approved",
        "Sensitive attributes still in the inputs"]),
    ("Future scope", [
        "Class weights / threshold tuning",
        "Drop sensitive features; fairness audit",
        "Try gradient boosting (XGBoost)",
        "Larger, real-world data"]),
]
for i, (head, pts) in enumerate(cols):
    x = 0.6 + i * 4.1
    card(s, x, 1.55, 3.85, 3.9, fill=GREEN_2)
    text(s, head, x + 0.3, 1.75, 3.3, 0.5, size=20, bold=True, color=GOLD)
    bullets(s, pts, x + 0.3, 2.4, 3.35, 3.0, size=16, color=WHITE)
text(s, "Thank you. Questions welcome.", 0.6, 5.95, 12.1, 0.6, size=24,
     bold=True, color=WHITE, font=HEAD_FONT, align=PP_ALIGN.CENTER)
finish(s, 14, "Mention a limitation before anyone asks; examiners reward "
              "self-awareness. Then thank the panel.", dark=True)

# ===========================================================================
prs.save(OUT_PATH)
print(f"Saved {len(prs.slides)} slides -> {OUT_PATH}")


def write_markdown_outline(path: str) -> None:
    """Text version of the deck (slide text + speaker notes), regenerated
    with the deck so the two can never disagree."""
    out = ["# Presentation: Loan Approval Prediction",
           "",
           "Text version of `Loan_Approval_Presentation.pptx`, generated by "
           "`src/build_presentation.py`. Do not edit by hand; rebuild instead.",
           ""]
    for number, slide in enumerate(prs.slides, start=1):
        texts, title = [], HEADINGS[slide.slide_id]
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t and t != f"{number} / {TOTAL}":
                    texts.append(t)
            elif shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                texts.append("\n".join(
                    ["| " + " | ".join(rows[0]) + " |",
                     "|" + "---|" * len(rows[0])]
                    + ["| " + " | ".join(r) + " |" for r in rows[1:]]))
            elif shape.has_chart:
                texts.append("*(chart: test-set metrics for each model)*")
            elif shape.shape_type == 13:   # picture
                texts.append(f"*(image: {shape.name})*")
        body = [t for t in texts if t != title]
        out += ["---", "", f"## Slide {number}: {title.splitlines()[0]}", ""]
        for t in body:
            out += [t if t.startswith(("|", "*(")) else
                    "\n".join(f"> {line}" if line else ">" for line in t.splitlines()), ""]
        notes = slide.notes_slide.notes_text_frame.text.strip()
        out += [f"**Say:** {notes}", ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    print(f"Saved text outline -> {path}")


write_markdown_outline(os.path.join(PROJECT_ROOT, "docs", "PRESENTATION_SLIDES.md"))
